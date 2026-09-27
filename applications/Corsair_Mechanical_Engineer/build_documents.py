from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT = "/home/user/Shom_brain_networks_AI/applications/Corsair_Mechanical_Engineer/"
NAME = "Arunabha Majumder"
CONTACT = "Aalborg, Denmark | +45 71 62 24 00 | somrkmv1997@gmail.com | linkedin.com/in/arunabha-majumder-681264107"


from reportlab.lib.pagesizes import A4
from reportlab.lib.units import cm
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, ListFlowable, ListItem, HRFlowable
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.fonts import addMapping
from xml.sax.saxutils import escape as esc
F = "/usr/share/fonts/truetype/liberation/LiberationSans-"
for n, f in (("Arial", "Regular"), ("Arial-B", "Bold"), ("Arial-I", "Italic"), ("Arial-BI", "BoldItalic")):
    pdfmetrics.registerFont(TTFont(n, F + f + ".ttf"))
addMapping("Arial", 0, 0, "Arial"); addMapping("Arial", 1, 0, "Arial-B")
addMapping("Arial", 0, 1, "Arial-I"); addMapping("Arial", 1, 1, "Arial-BI")

def new_doc():
    d = Document()
    d.story = []
    s = d.sections[0]
    s.page_height, s.page_width = Cm(29.7), Cm(21.0)
    s.left_margin = s.right_margin = Cm(1.9)
    s.top_margin = s.bottom_margin = Cm(1.6)
    st = d.styles["Normal"]
    st.font.name = "Arial"; st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    st.paragraph_format.space_after = Pt(2)
    st.paragraph_format.line_spacing = 1.08
    return d

def rl_style(size=10.5, after=2, before=0):
    return ParagraphStyle("x", fontName="Arial", fontSize=size, leading=size * 1.25,
                          spaceAfter=after, spaceBefore=before)

def para(d, text="", bold=False, size=None, align=None, after=None, before=None, italic=False):
    p = d.add_paragraph()
    if text:
        r = p.add_run(text); r.bold = bold; r.italic = italic
        if size: r.font.size = Pt(size)
    if align: p.alignment = align
    if after is not None: p.paragraph_format.space_after = Pt(after)
    if before is not None: p.paragraph_format.space_before = Pt(before)
    p.rl = [text, bold, italic, size or 10.5, 2 if after is None else after, before or 0]
    d.story.append(p)
    return p

def rule(p):
    pPr = p._p.get_or_add_pPr()
    b = OxmlElement("w:pBdr"); bt = OxmlElement("w:bottom")
    for k, v in (("val", "single"), ("sz", "6"), ("space", "1"), ("color", "404040")):
        bt.set(qn("w:" + k), v)
    b.append(bt); pPr.append(b)
    p.rule = True

def header(d, title):
    para(d, NAME, bold=True, size=18, after=0)
    para(d, title, bold=True, size=11, after=0)
    para(d, CONTACT, size=9.5, after=6)

def heading(d, text):
    p = para(d, text.upper(), bold=True, size=11, before=8, after=3)
    rule(p)

def bullet(d, text, lead=None):
    p = d.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2)
    if lead:
        p.add_run(lead).bold = True
    p.add_run(text)
    d.story.append(("bullet", lead, text))

def role(d, title, dates, org):
    p = d.add_paragraph()
    p.paragraph_format.space_after = Pt(0); p.paragraph_format.space_before = Pt(4)
    p.add_run(title).bold = True
    p.add_run(" | " + dates)
    d.story.append(("role", title, dates))
    para(d, org, italic=True, after=2)

def to_pdf(d, path, top=1.6):
    fl = []
    for it in d.story:
        if isinstance(it, tuple) and it[0] == "bullet":
            txt = ("<b>%s</b>" % esc(it[1]) if it[1] else "") + esc(it[2])
            fl.append(ListFlowable([ListItem(Paragraph(txt, rl_style(after=0)), leftIndent=12)],
                                   bulletType="bullet", start="•", leftIndent=12, bulletFontSize=9))
            fl.append(Spacer(0, 2))
        elif isinstance(it, tuple) and it[0] == "role":
            fl.append(Paragraph("<b>%s</b> | %s" % (esc(it[1]), esc(it[2])), rl_style(after=0, before=4)))
        else:
            text, bold, italic, size, after, before = it.rl
            t = esc(text)
            if bold: t = "<b>%s</b>" % t
            if italic: t = "<i>%s</i>" % t
            fl.append(Paragraph(t or "&nbsp;", rl_style(size, 0 if getattr(it, "rule", False) else after, before)))
            if getattr(it, "rule", False):
                fl.append(HRFlowable(width="100%", thickness=0.6, color="#404040", spaceBefore=1, spaceAfter=after))
    SimpleDocTemplate(path, pagesize=A4, leftMargin=1.9 * cm, rightMargin=1.9 * cm, topMargin=top * cm,
                      bottomMargin=1.6 * cm, title=path.split("/")[-1][:-4].replace("_", " "),
                      author=NAME).build(fl)

# ---------------- RESUME ----------------
d = new_doc()
header(d, "Mechanical Engineer | New Product Development | Mechatronic Systems | SolidWorks")

heading(d, "Professional Summary")
para(d, "Mechanical Engineer with a Master's in Mechatronics and about three years of hands-on new product "
        "development at Aalborg University, taking electromechanical devices from concept through SolidWorks CAD, "
        "FEA, prototype fabrication, assembly and testing. Designed a hybrid-actuated shoulder exoskeleton built from "
        "machined and additively manufactured parts that won the Best Paper Award at IFToMM ISRM 2026. Certified "
        "SolidWorks Associate (CSWA) with a structured, analytical approach to problem solving and experience "
        "collaborating across mechanical, electronics, controls and biomechanics disciplines in international "
        "teams in Denmark and India.")

heading(d, "Core Skills")
bullet(d, "SolidWorks (CSWA certified), Fusion 360, Finite Element Analysis (FEA), MATLAB/Simulink",
       "CAD & Simulation: ")
bullet(d, "New Product Development (NPD), Mechanical Design, Design for Assembly (DFA), Additive Manufacturing / "
          "3D Printing, Machined Parts, Prototype Fabrication, Assembly, Custom Test Rigs", "Design & Manufacturing: ")
bullet(d, "Actuator Design, Motor / Encoder / Force-Sensor Integration, DAQ, Embedded Control, Pneumatic Systems, "
          "Control Systems", "Mechatronic Systems: ")
bullet(d, "Stiffness and Kinematic Modelling, Experimental Validation, Test Planning, Data Analysis, "
          "Motion Capture (Qualisys), EMG", "Analysis & Testing: ")
bullet(d, "English (C1, fluent written and spoken), cross-disciplinary collaboration, ownership from concept "
          "to tested hardware", "Communication: ")

heading(d, "Professional Experience")
role(d, "Research Engineer – Robotics & Actuator Development (PhD Project: VIEXO)", "2023 – Present",
     "Aalborg University, Aalborg, Denmark")
bullet(d, "Owned the full product development cycle of a wearable robotic device from concept to tested hardware: "
          "SolidWorks CAD and FEA, mechanical fabrication, assembly, motor/encoder/force-sensor and DAQ "
          "integration, embedded control and custom automated test rigs.")
bullet(d, "Cut shoulder exoskeleton energy consumption by more than 25% by designing a hybrid actuation architecture "
          "that combines a parallel spring with a 6 Nm motor to offload gravity torque; awarded Best Paper, "
          "IFToMM ISRM 2026.")
bullet(d, "Removed the need for an external force/torque sensor in a 5-bar parallel robot by deriving a sensorless "
          "force-estimation method from the actuator torque–deflection characteristic, reducing system cost, mass "
          "and integration complexity.")
bullet(d, "Validated analytical stiffness models against experimental hardware to below 10% RMS error by mapping "
          "variable-stiffness actuator (VSA) states into Cartesian workspace stiffness, de-risking the mechanism "
          "before fabrication.")
bullet(d, "Reduced user anterior deltoid muscle effort by more than 15% by deploying an Assist-as-Needed hybrid "
          "controller on functional prototype hardware, validated on human subjects.")
bullet(d, "Planned and ran a 12-participant validation study across 3 assistance modes and 2 load conditions, "
          "covering instrumentation (EMG, motion capture), ethics protocol and the end-to-end data pipeline; "
          "results support two journal submissions.")

role(d, "M.Tech Researcher – Actuator Systems (Master's Thesis)", "2020 – 2022",
     "CSIR-Central Mechanical Engineering Research Institute (CMERI), Durgapur, India")
bullet(d, "Reduced dynamic position-tracking error to 0.3–0.78% by designing a neural-network-based gain-scheduled "
          "PID controller for a non-linear pneumatic actuator, outperforming the classical PID baseline under "
          "varying axial load.")
bullet(d, "Fabricated pneumatic artificial muscle actuators in-house and integrated valves, pressure sensors and DAQ "
          "to characterise load–displacement behaviour and derive an empirical actuator model.")

role(d, "Mechanical Engineering Capstone – Upper-Limb Rehabilitation Exoskeleton", "2018 – 2019",
     "Siddaganga Institute of Technology, Tumakuru, India")
bullet(d, "Led CAD design, fabrication and user testing of an upper-limb rehabilitation exoskeleton; won the Best "
          "Major Project Award and an INR 100,000 TATA Technologies Innovation Grant. "
          "Demo: youtube.com/watch?v=q5Ystz9NCpQ")

heading(d, "Education")
role(d, "Ph.D. Candidate, Mechatronics & Robotics", "2023 – Present",
     "Aalborg University, Department of Materials and Production, Aalborg, Denmark")
role(d, "M.Tech, Mechatronics", "2020 – 2022",
     "AcSIR – CSIR-Central Mechanical Engineering Research Institute (CMERI), Durgapur, India")
role(d, "B.Tech, Mechanical Engineering", "2015 – 2019",
     "Siddaganga Institute of Technology, Visvesvaraya Technological University (VTU), Tumakuru, India")

heading(d, "Engineering Projects")
bullet(d, "Led end-to-end design and fabrication of a custom Go-Kart built from the ground up.", "Go-Kart (2017): ")
bullet(d, "Worked in a multidisciplinary team to engineer a human-electric hybrid vehicle for the Efficycle "
          "competition.", "Efficycle (2018): ")

heading(d, "Certifications & Awards")
bullet(d, "Certified SolidWorks Associate (CSWA) – Mechanical Design")
bullet(d, "Best Research Paper Award, 9th IFToMM International Symposium on Robotics and Mechatronics (2026)")
bullet(d, "Best Major Project Award, Department of Mechanical Engineering (2019)")
bullet(d, "TATA Technologies Innovation Award – INR 100,000 Grant (2019)")

heading(d, "Publications")
bullet(d, "Majumder, A., Wagner, J.W., Zhu, Y., Oliveira, A.S., and Bai, S. \"A Hybrid Actuated Shoulder Exoskeleton "
          "for Energy-Efficient Upper Arm Support.\" Robotics and Mechatronics: ISRM 2026, MMS, vol. 213, Springer. "
          "(Best Research Paper Award)")
bullet(d, "Majumder, A., Sarkar, D., Chakraborty, S., Singh, A., Roy, S.S., and Arora, A. \"Neural Network-Based Gain "
          "Scheduled Position Control of a Pneumatic Artificial Muscle.\" IEEE CONECCT, 2022.")

heading(d, "Technical Tools")
bullet(d, "MATLAB/Simulink, Python, LabVIEW (basic), Arduino, LaTeX", "Software & Programming: ")
bullet(d, "SolidWorks, Fusion 360, FEA, Bambu Lab and Ultimaker Cura (3D printing)", "CAD & Manufacturing: ")
bullet(d, "Qualisys Motion Capture, Forsentec Force/Torque Sensors, Pneumatic Control, Custom Rig Fabrication",
       "Instrumentation & Testing: ")
bullet(d, "Adobe Illustrator, Adobe Photoshop; Diploma in Fine Arts (Painting), Bangiya Sangeet Parishad – "
          "sketching and technical illustration", "Visual Design: ")

heading(d, "Languages")
para(d, "English (C1, fluent) | Bengali (Native) | Hindi (Conversational)")

heading(d, "References")
para(d, "Available upon request.")
d.save(OUT + "Arunabha_Majumder_Resume_Corsair_Mechanical_Engineer.docx")
to_pdf(d, OUT + "Arunabha_Majumder_Resume_Corsair_Mechanical_Engineer.pdf")

# ---------------- COVER LETTER ----------------
c = new_doc()
c.sections[0].top_margin = Cm(2.0)
header(c, "Mechanical Engineer | New Product Development | Mechatronic Systems")
para(c, "September 27, 2026", after=8)
para(c, "Hiring Team, Mechanical Engineering", after=0)
para(c, "Corsair", after=0)
para(c, "Landshut, Bavaria, Germany", after=8)
para(c, "Application: Mechanical Engineer (On-site, Landshut)", bold=True, after=8)
para(c, "Dear Hiring Team,", after=8)
body = [
 "I am applying for the Mechanical Engineer position at Corsair in Landshut. I am a mechanical engineer with a "
 "Master's in Mechatronics, finishing my PhD in Mechatronics and Robotics at Aalborg University. For the past three "
 "years I have developed electromechanical products from first concept to tested hardware, and new product "
 "development is exactly the work I want to do.",

 "In my PhD, I led the mechanical development of a hybrid-actuated shoulder exoskeleton. I created the SolidWorks "
 "models and drawings, iterated the design on FEA results for stiffness, strength and weight, built the parts "
 "through machining and additive manufacturing, and assembled and tested the prototypes. I also integrated the "
 "motor, encoder, force sensors and DAQ with embedded control, and built custom test rigs to validate the design. "
 "The hybrid actuation cut energy consumption by more than 25% and won the Best Paper Award at IFToMM ISRM 2026.",

 "I take a structured, analytical approach to engineering problems. I validated my analytical stiffness models "
 "against physical tests to within 10% RMS error before committing to fabrication. On a 5-bar parallel robot, I "
 "removed an external force/torque sensor by estimating force from the actuator itself, which reduced cost, mass "
 "and integration complexity. During my Master's at CSIR-CMERI, I fabricated pneumatic actuators in-house and "
 "developed a neural-network-based controller that reduced position-tracking error to below 0.8%.",

 "My work sits between mechanical design, electronics, controls and biomechanics, so collaborating across "
 "disciplines is part of my daily routine. I have worked in research environments in both India and Denmark, with "
 "English as my working language, and I am a Certified SolidWorks Associate.",

 "My hands-on manufacturing experience so far is mainly with machined and additively manufactured parts. Designing "
 "for injection moulding, metal casting and sheet metal in series production is where I want to grow next, and "
 "Corsair's product development team is the right place to do it.",

 "I would welcome the opportunity to discuss how I can contribute to Corsair. References are available upon request.",
]
for t in body:
    p = para(c, t, after=8); p.alignment = WD_ALIGN_PARAGRAPH.LEFT
para(c, "Kind regards,", after=2)
para(c, NAME)
c.save(OUT + "Arunabha_Majumder_Cover_Letter_Corsair_Mechanical_Engineer.docx")
to_pdf(c, OUT + "Arunabha_Majumder_Cover_Letter_Corsair_Mechanical_Engineer.pdf", top=2.0)
