"""Shared ATS-friendly DOCX + PDF layout helpers for application documents."""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

NAME = "Arunabha Majumder"
CONTACT = "Aalborg, Denmark | +45 71 62 24 00 | somrkmv1997@gmail.com | linkedin.com/in/arunabha-majumder-681264107"
LINKS = "Portfolio: arunabha22.github.io/arunabha-majumder | Project: www.viexo.aau.dk"


from reportlab.lib.pagesizes import A4
from reportlab.lib.units import cm
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, ListFlowable, ListItem, HRFlowable
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.fonts import addMapping
from xml.sax.saxutils import escape as esc
F = "/usr/share/fonts/truetype/liberation/LiberationSans-"
for n, f in (("Arial", "Regular"), ("Arial-B", "Bold"), ("Arial-I", "Italic"), ("Arial-BI", "BoldItalic")):
    pdfmetrics.registerFont(TTFont(n, F + f + ".ttf"))
addMapping("Arial", 0, 0, "Arial"); addMapping("Arial", 1, 0, "Arial-B")
addMapping("Arial", 0, 1, "Arial-I"); addMapping("Arial", 1, 1, "Arial-BI")

# Plain, human typography: no em dashes, no invisible/zero-width characters.
_TIDY = [(" \u2014 ", " | "), ("\u2014", " - "), ("\u200b", ""), ("\u200c", ""), ("\u200d", ""),
         ("\ufeff", ""), ("\u00a0", " "), ("\u2018", "'"), ("\u2019", "'"), ("\u201c", '"'), ("\u201d", '"')]

def tidy(t):
    for a, b in _TIDY:
        t = t.replace(a, b)
    return t

def new_doc():
    d = Document()
    cp = d.core_properties  # no tool watermark in file properties
    cp.author = cp.last_modified_by = NAME
    cp.comments = cp.keywords = cp.subject = cp.category = ""
    d.story = []
    s = d.sections[0]
    s.page_height, s.page_width = Cm(29.7), Cm(21.0)
    s.left_margin = s.right_margin = Cm(1.9)
    s.top_margin = s.bottom_margin = Cm(1.6)
    st = d.styles["Normal"]
    st.font.name = "Arial"; st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    st.paragraph_format.space_after = Pt(2)
    st.paragraph_format.line_spacing = 1.08
    return d

def rl_style(size=10.5, after=2, before=0):
    return ParagraphStyle("x", fontName="Arial", fontSize=size, leading=size * 1.25,
                          spaceAfter=after, spaceBefore=before)

def para(d, text="", bold=False, size=None, align=None, after=None, before=None, italic=False):
    text = tidy(text)
    p = d.add_paragraph()
    if text:
        r = p.add_run(text); r.bold = bold; r.italic = italic
        if size: r.font.size = Pt(size)
    if align: p.alignment = align
    if after is not None: p.paragraph_format.space_after = Pt(after)
    if before is not None: p.paragraph_format.space_before = Pt(before)
    p.rl = [text, bold, italic, size or 10.5, 2 if after is None else after, before or 0]
    d.story.append(p)
    return p

def rule(p):
    pPr = p._p.get_or_add_pPr()
    b = OxmlElement("w:pBdr"); bt = OxmlElement("w:bottom")
    for k, v in (("val", "single"), ("sz", "6"), ("space", "1"), ("color", "404040")):
        bt.set(qn("w:" + k), v)
    b.append(bt); pPr.append(b)
    p.rule = True

def header(d, title, links=False):
    para(d, NAME, bold=True, size=18, after=0)
    para(d, title, bold=True, size=11, after=0)
    para(d, CONTACT, size=9.5, after=0 if links else 6)
    if links:
        para(d, LINKS, size=9.5, after=6)

def heading(d, text):
    p = para(d, text.upper(), bold=True, size=11, before=8, after=3)
    rule(p)

def bullet(d, text, lead=None):
    text, lead = tidy(text), tidy(lead) if lead else lead
    p = d.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2)
    if lead:
        p.add_run(lead).bold = True
    p.add_run(text)
    d.story.append(("bullet", lead, text))

def role(d, title, dates, org):
    title, dates = tidy(title), tidy(dates)
    p = d.add_paragraph()
    p.paragraph_format.space_after = Pt(0); p.paragraph_format.space_before = Pt(4)
    p.add_run(title).bold = True
    if dates:
        p.add_run(" | " + dates)
    d.story.append(("role", title, dates))
    para(d, org, italic=True, after=2)

def to_pdf(d, path, top=1.6):
    fl = []
    for it in d.story:
        if isinstance(it, tuple) and it[0] == "bullet":
            txt = ("<b>%s</b>" % esc(it[1]) if it[1] else "") + esc(it[2])
            fl.append(ListFlowable([ListItem(Paragraph(txt, rl_style(after=0)), leftIndent=12)],
                                   bulletType="bullet", start="•", leftIndent=12, bulletFontSize=9))
            fl.append(Spacer(0, 2))
        elif isinstance(it, tuple) and it[0] == "role":
            fl.append(Paragraph("<b>%s</b>" % esc(it[1]) + (" | %s" % esc(it[2]) if it[2] else ""), rl_style(after=0, before=4)))
        else:
            text, bold, italic, size, after, before = it.rl
            t = esc(text)
            if bold: t = "<b>%s</b>" % t
            if italic: t = "<i>%s</i>" % t
            fl.append(Paragraph(t or "&nbsp;", rl_style(size, 0 if getattr(it, "rule", False) else after, before)))
            if getattr(it, "rule", False):
                fl.append(HRFlowable(width="100%", thickness=0.6, color="#404040", spaceBefore=1, spaceAfter=after))
    SimpleDocTemplate(path, pagesize=A4, leftMargin=1.9 * cm, rightMargin=1.9 * cm, topMargin=top * cm,
                      bottomMargin=1.6 * cm, title=path.split("/")[-1][:-4].replace("_", " "),
                      author=NAME, creator=NAME, producer=NAME, subject="", keywords="").build(fl)

